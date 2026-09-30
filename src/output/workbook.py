"""
Multi-sheet Excel workbook — spec §33.

Workbook structure (§33.1):
    Data           — one row per record, one column per field
    Changes        — new/modified/removed records with color coding (§33.2)
    Sources        — per-field provenance
    Quality        — quality report per rule
    Confidence     — per-field confidence when <field>_confidence keys exist
    Run Summary    — metadata, cost, run IDs
    Errors & Warnings — non-fatal issues

Windows-safe atomic write with retry-on-PermissionError (same pattern as
src/diff/excel_writer.py).
"""

from __future__ import annotations
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.trust.cleaning import MISSING


# ---------------------------------------------------------------------------
# Colour palette (§33.2 — accessibility-safe)
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

_NEW_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_MODIFIED_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_REMOVED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


# ---------------------------------------------------------------------------
# Build result
# ---------------------------------------------------------------------------

@dataclass
class WorkbookResult:
    path: Path
    sheet_names: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "path": str(self.path),
            "sheet_names": list(self.sheet_names),
            "warnings": list(self.warnings),
        }


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _clean_cell(value: Any) -> Any:
    if value is MISSING or value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, dict):
        import json as _json
        return _json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, (list, tuple)):
        return ", ".join(_clean_cell(v) or "" for v in value)
    return str(value)


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center")


def _autofit(ws, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        try:
            length = max(
                len(str(c.value)) if c.value is not None else 0
                for c in column_cells
            )
        except ValueError:
            continue
        letter = get_column_letter(column_cells[0].column)
        ws.column_dimensions[letter].width = min(max_width, max(10, length + 3))


def _write_kv_sheet(ws, rows: list[tuple[str, Any]]) -> None:
    ws.append(["Key", "Value"])
    _style_header(ws)
    for k, v in rows:
        ws.append([k, _clean_cell(v)])
    _autofit(ws)


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_data_sheet(wb: Workbook, records: list[dict]) -> None:
    ws = wb.create_sheet("Data")
    if not records:
        ws.append(["(no records)"])
        return

    columns: list[str] = []
    seen: set[str] = set()
    for r in records:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                columns.append(k)

    ws.append(columns)
    _style_header(ws)

    for r in records:
        ws.append([_clean_cell(r.get(c)) for c in columns])
    _autofit(ws)


def _build_changes_sheet(wb: Workbook, change_set) -> Optional[str]:
    if change_set is None:
        return None
    ws = wb.create_sheet("Changes")
    ws.append(["Status", "Identity", "Field", "Old", "New", "Record Summary"])
    _style_header(ws)

    def _summary(rec: dict) -> str:
        parts = []
        for k in ("title", "business_name", "name"):
            if rec.get(k):
                parts.append(str(rec[k]))
        return " | ".join(parts) if parts else "(no summary)"

    for rc in change_set.new:
        ws.append(["NEW", rc.identity, "", "", "", _summary(rc.record)])
        for cell in ws[ws.max_row]:
            cell.fill = _NEW_FILL

    for rc in change_set.removed:
        ws.append(["REMOVED", rc.identity, "", "", "", _summary(rc.record)])
        for cell in ws[ws.max_row]:
            cell.fill = _REMOVED_FILL

    for rc in change_set.modified:
        if not rc.field_changes:
            ws.append(["MODIFIED", rc.identity, "(no field detail)", "", "",
                       _summary(rc.record)])
            for cell in ws[ws.max_row]:
                cell.fill = _MODIFIED_FILL
            continue
        for fc in rc.field_changes:
            ws.append(["MODIFIED", rc.identity, fc.field_name,
                       _clean_cell(fc.old_value), _clean_cell(fc.new_value),
                       _summary(rc.record)])
            for cell in ws[ws.max_row]:
                cell.fill = _MODIFIED_FILL

    _autofit(ws)
    return "Changes"


def _build_sources_sheet(wb: Workbook, provenance_records: list) -> Optional[str]:
    if not provenance_records:
        return None
    ws = wb.create_sheet("Sources")
    headers = [
        "record_id", "field", "source_url", "domain",
        "observed_at", "fetched_at", "method",
        "raw_value", "normalized_value", "confidence",
    ]
    ws.append(headers)
    _style_header(ws)

    for p in provenance_records:
        # Support both ProvenanceRecord objects and plain dicts
        if hasattr(p, "to_dict"):
            d = p.to_dict()
        else:
            d = p
        ws.append([
            d.get("record_id"), d.get("field_name"), d.get("source_url"),
            d.get("source_domain"), d.get("observed_at"), d.get("fetched_at"),
            d.get("extraction_method"), d.get("raw_value"),
            d.get("normalized_value"), d.get("confidence"),
        ])
    _autofit(ws)
    return "Sources"


def _build_quality_sheet(wb: Workbook, quality_result) -> Optional[str]:
    if quality_result is None:
        return None
    ws = wb.create_sheet("Quality")
    ws.append(["Field", "Value"])
    _style_header(ws)

    ws.append(["passed", "YES" if quality_result.passed else "NO"])
    ws.append(["score", quality_result.score])
    ws.append(["suggested_action", quality_result.suggested_action])
    ws.append(["failed_rules", ", ".join(quality_result.failed_rules) or "(none)"])
    ws.append(["explanation", quality_result.explanation])
    ws.append(["", ""])
    ws.append(["Rule", "Passed", "Detail", "Value"])
    _style_header(ws, row=ws.max_row)

    for rr in quality_result.per_rule_results:
        ws.append([
            rr.name,
            "PASS" if rr.passed else "FAIL",
            rr.detail,
            _clean_cell(rr.value),
        ])
        fill = _PASS_FILL if rr.passed else _FAIL_FILL
        for cell in ws[ws.max_row]:
            cell.fill = fill

    _autofit(ws)
    return "Quality"


def _build_confidence_sheet(wb: Workbook, records: list[dict]) -> Optional[str]:
    """
    Only build this sheet if at least one record has a `<field>_confidence`
    or `confidence` key. Otherwise skip.
    """
    confidence_columns: set[str] = set()
    for r in records:
        for k in r.keys():
            if k == "confidence" or k.endswith("_confidence"):
                confidence_columns.add(k)

    if not confidence_columns:
        return None

    ws = wb.create_sheet("Confidence")
    identity_key = None
    for candidate in ("title", "business_name", "name", "url"):
        if records and any(candidate in r for r in records):
            identity_key = candidate
            break

    headers = ([identity_key] if identity_key else ["record_index"]) + sorted(confidence_columns)
    ws.append(headers)
    _style_header(ws)

    for i, r in enumerate(records):
        ident = r.get(identity_key) if identity_key else i
        row = [ident] + [_clean_cell(r.get(c)) for c in sorted(confidence_columns)]
        ws.append(row)

    _autofit(ws)
    return "Confidence"


def _build_run_summary_sheet(wb: Workbook, metadata: dict) -> str:
    ws = wb.create_sheet("Run Summary")
    _write_kv_sheet(ws, sorted(metadata.items()))
    return "Run Summary"


def _build_errors_sheet(wb: Workbook, errors: list[str]) -> Optional[str]:
    if not errors:
        return None
    ws = wb.create_sheet("Errors & Warnings")
    ws.append(["Message"])
    _style_header(ws)
    for e in errors:
        ws.append([str(e)])
    _autofit(ws)
    return "Errors & Warnings"


# ---------------------------------------------------------------------------
# Orchestrator
# ---------------------------------------------------------------------------

class WorkbookBuilder:
    def __init__(self, retries: int = 3):
        self.retries = retries

    def build(
        self,
        output_path: str | Path,
        records: list[dict],
        change_set=None,
        quality_result=None,
        provenance_records: Optional[list] = None,
        run_metadata: Optional[dict] = None,
        errors: Optional[list[str]] = None,
    ) -> WorkbookResult:
        path = Path(output_path)
        warnings: list[str] = []

        # Remove default sheet — we build them ourselves.
        wb = Workbook()
        wb.remove(wb.active)

        _build_data_sheet(wb, records)

        if _build_changes_sheet(wb, change_set) is None:
            pass
        if _build_sources_sheet(wb, provenance_records or []) is None:
            pass
        if _build_quality_sheet(wb, quality_result) is None:
            pass
        if _build_confidence_sheet(wb, records) is None:
            pass

        _build_run_summary_sheet(wb, run_metadata or {})
        _build_errors_sheet(wb, errors or [])

        sheet_names = wb.sheetnames

        # Atomic write with retry on Windows file-lock
        tmp_path = path.with_suffix(".tmp.xlsx")
        wb.save(tmp_path)

        delay = 2
        for attempt in range(1, self.retries + 1):
            try:
                tmp_path.replace(path)
                return WorkbookResult(path=path, sheet_names=sheet_names,
                                      warnings=warnings)
            except PermissionError:
                warnings.append(
                    f"file locked (attempt {attempt}); retrying in {delay}s"
                )
                time.sleep(delay)
                delay *= 2

        # If we exhausted retries, leave the tmp file and warn.
        warnings.append(
            f"could not replace {path} after {self.retries} attempts; "
            f"wrote to {tmp_path} instead"
        )
        return WorkbookResult(path=tmp_path, sheet_names=sheet_names,
                              warnings=warnings)


# ---------------------------------------------------------------------------
# Convenience
# ---------------------------------------------------------------------------

def build_workbook(
    output_path: str | Path,
    records: list[dict],
    **kwargs,
) -> WorkbookResult:
    return WorkbookBuilder().build(output_path, records, **kwargs)


# ---------------------------------------------------------------------------
# Smoke test
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    import tempfile
    from src.history.change import classify
    from src.quality.rules import QualityResult, RuleResult
    from src.trust.provenance import ProvenanceStore

    records = [
        {"title": "A", "price": "$10", "confidence": 0.9},
        {"title": "B", "price": "$20", "confidence": 0.7},
    ]
    prev = [{"title": "A", "price": "$9"}]

    change_set = classify(prev, records)
    quality = QualityResult(
        passed=True, score=1.0,
        per_rule_results=[RuleResult("min_records", True, "ok", 2)],
    )
    prov = ProvenanceStore(run_id="r", job_id="j")
    prov.record_batch(records, source_url="https://x.com/a", extraction_method="css")

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        r = build_workbook(
            out, records,
            change_set=change_set,
            quality_result=quality,
            provenance_records=prov.to_list(),
            run_metadata={"task": "demo", "run_id": "r", "pages": 3},
            errors=["page 4 timed out"],
        )
        assert out.exists(), r.warnings
        assert "Data" in r.sheet_names
        assert "Changes" in r.sheet_names
        assert "Sources" in r.sheet_names
        assert "Quality" in r.sheet_names
        assert "Confidence" in r.sheet_names
        assert "Run Summary" in r.sheet_names
        assert "Errors & Warnings" in r.sheet_names
        print("sheets:", r.sheet_names)

        # Without optional inputs, those sheets are skipped
        out2 = Path(tmp) / "minimal.xlsx"
        r2 = build_workbook(out2, records)
        assert "Data" in r2.sheet_names
        assert "Run Summary" in r2.sheet_names
        assert "Changes" not in r2.sheet_names
        assert "Sources" not in r2.sheet_names
        assert "Quality" not in r2.sheet_names
        assert "Confidence" in r2.sheet_names   # confidence present on records
        assert "Errors & Warnings" not in r2.sheet_names
        print("minimal sheets:", r2.sheet_names)

    print("WorkbookBuilder OK.")