"""Unit tests for multi-sheet Excel workbook (spec §33)."""
from pathlib import Path

from openpyxl import load_workbook

from src.output.workbook import build_workbook, WorkbookBuilder
from src.history.change import classify
from src.quality.rules import QualityResult, RuleResult
from src.trust.provenance import ProvenanceStore


SAMPLE = [
    {"title": "A", "price": "$10", "confidence": 0.9},
    {"title": "B", "price": "$20", "confidence": 0.7},
]


def _read_sheets(path: Path) -> dict[str, list[list]]:
    wb = load_workbook(path)
    out: dict[str, list[list]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = [list(row) for row in ws.iter_rows(values_only=True)]
    return out


# --- basic structure --------------------------------------------------

def test_build_creates_file(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    r = build_workbook(out, SAMPLE)
    assert out.exists()
    assert r.path == out


def test_data_sheet_always_present(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    build_workbook(out, SAMPLE)
    sheets = _read_sheets(out)
    assert "Data" in sheets
    data = sheets["Data"]
    assert data[0] == ["title", "price", "confidence"]
    assert data[1] == ["A", "$10", 0.9]


def test_run_summary_always_present(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    build_workbook(out, SAMPLE, run_metadata={"task": "demo"})
    sheets = _read_sheets(out)
    assert "Run Summary" in sheets
    summary = {row[0]: row[1] for row in sheets["Run Summary"][1:]}
    assert summary["task"] == "demo"


# --- optional sheets appear only when data present --------------------

def test_changes_sheet_only_when_change_set(tmp_path: Path):
    # Without change set
    out = tmp_path / "a.xlsx"
    build_workbook(out, SAMPLE)
    sheets = _read_sheets(out)
    assert "Changes" not in sheets

    # With change set
    prev = [{"title": "A", "price": "$9"}]
    cs = classify(prev, SAMPLE)
    out = tmp_path / "b.xlsx"
    build_workbook(out, SAMPLE, change_set=cs)
    sheets = _read_sheets(out)
    assert "Changes" in sheets


def test_changes_sheet_colour_codes_new_modified_removed(tmp_path: Path):
    prev = [
        {"title": "A", "price": "$9"},
        {"title": "C", "price": "$30"},
    ]
    curr = [
        {"title": "A", "price": "$10"},   # MODIFIED
        {"title": "B", "price": "$20"},   # NEW
        # C removed
    ]
    cs = classify(prev, curr)
    out = tmp_path / "c.xlsx"
    build_workbook(out, curr, change_set=cs)

    wb = load_workbook(out)
    ws = wb["Changes"]
    statuses = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert "NEW" in statuses
    assert "MODIFIED" in statuses
    assert "REMOVED" in statuses


def test_sources_sheet_only_when_provenance(tmp_path: Path):
    out = tmp_path / "a.xlsx"
    build_workbook(out, SAMPLE)
    assert "Sources" not in _read_sheets(out)

    prov = ProvenanceStore(run_id="r", job_id="j")
    prov.record_batch(SAMPLE, source_url="https://x.com/a", extraction_method="css")
    out = tmp_path / "b.xlsx"
    build_workbook(out, SAMPLE, provenance_records=prov.to_list())
    sheets = _read_sheets(out)
    assert "Sources" in sheets
    header = sheets["Sources"][0]
    assert "source_url" in header
    assert "field" in header


def test_quality_sheet_only_when_quality_result(tmp_path: Path):
    out = tmp_path / "a.xlsx"
    build_workbook(out, SAMPLE)
    assert "Quality" not in _read_sheets(out)

    q = QualityResult(
        passed=True, score=1.0,
        per_rule_results=[RuleResult("min_records", True, "2 >= 1", 2)],
    )
    out = tmp_path / "b.xlsx"
    build_workbook(out, SAMPLE, quality_result=q)
    sheets = _read_sheets(out)
    assert "Quality" in sheets
    # Check the per-rule row exists
    all_text = "\n".join(str(c) for row in sheets["Quality"] for c in row if c is not None)
    assert "min_records" in all_text


def test_confidence_sheet_only_when_confidence_columns(tmp_path: Path):
    records_without = [{"title": "A", "price": "$10"}]
    out = tmp_path / "a.xlsx"
    build_workbook(out, records_without)
    assert "Confidence" not in _read_sheets(out)

    out = tmp_path / "b.xlsx"
    build_workbook(out, SAMPLE)
    assert "Confidence" in _read_sheets(out)


def test_errors_sheet_only_when_errors(tmp_path: Path):
    out = tmp_path / "a.xlsx"
    build_workbook(out, SAMPLE)
    assert "Errors & Warnings" not in _read_sheets(out)

    out = tmp_path / "b.xlsx"
    build_workbook(out, SAMPLE, errors=["page 4 failed"])
    sheets = _read_sheets(out)
    assert "Errors & Warnings" in sheets
    assert any("page 4 failed" in str(c) for row in sheets["Errors & Warnings"] for c in row)


# --- record edge cases ------------------------------------------------

def test_empty_records_still_produces_file(tmp_path: Path):
    out = tmp_path / "e.xlsx"
    r = build_workbook(out, [])
    assert out.exists()
    assert "Data" in r.sheet_names


def test_missing_sentinel_written_as_blank(tmp_path: Path):
    from src.trust.cleaning import MISSING
    out = tmp_path / "s.xlsx"
    build_workbook(out, [{"a": MISSING, "b": "ok"}])
    sheets = _read_sheets(out)
    # Row 2: a is None (blank), b is "ok"
    assert sheets["Data"][1][0] is None
    assert sheets["Data"][1][1] == "ok"


# --- result object ----------------------------------------------------

def test_result_lists_sheet_names(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    r = build_workbook(out, SAMPLE)
    assert "Data" in r.sheet_names
    assert "Run Summary" in r.sheet_names


def test_result_to_dict(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    r = build_workbook(out, SAMPLE)
    d = r.to_dict()
    assert "path" in d
    assert "sheet_names" in d
    assert "warnings" in d