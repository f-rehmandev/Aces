"""Verify diff_status renders as a column in the workbook when present."""
from pathlib import Path

from openpyxl import load_workbook

from src.output.workbook import build_workbook


def _read_sheets(path: Path) -> dict[str, list[list]]:
    wb = load_workbook(path)
    return {
        name: [list(row) for row in wb[name].iter_rows(values_only=True)]
        for name in wb.sheetnames
    }


def test_diff_status_column_present_when_present(tmp_path: Path):
    records = [
        {"business_name": "A", "phone": "0300", "diff_status": "NEW"},
        {"business_name": "B", "phone": "0301", "diff_status": "EXISTING"},
        {"business_name": "C", "phone": "0302", "diff_status": "NEW_UNVERIFIED"},
    ]
    out = tmp_path / "leads.xlsx"
    build_workbook(out, records)

    sheets = _read_sheets(out)
    header = sheets["Data"][0]
    assert "diff_status" in header

    idx = header.index("diff_status")
    values = [row[idx] for row in sheets["Data"][1:]]
    assert values == ["NEW", "EXISTING", "NEW_UNVERIFIED"]


def test_diff_status_absent_when_not_on_records(tmp_path: Path):
    records = [{"business_name": "A", "phone": "0300"}]
    out = tmp_path / "no_status.xlsx"
    build_workbook(out, records)

    sheets = _read_sheets(out)
    assert "diff_status" not in sheets["Data"][0]